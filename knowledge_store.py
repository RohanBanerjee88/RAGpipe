"""Versioned collections of source-preserving passages and dataset cards."""

import csv
import hashlib
import itertools
import re
from collections import deque
from pathlib import Path
from urllib.parse import urldefrag, urljoin, urlsplit, urlunsplit
from urllib.robotparser import RobotFileParser

from filelock import FileLock

from config import FAQ_JSON_PATH
from ingestion import sha256_text, utc_now_iso
from local_store import data_root, read_json, write_json


SUPPORTED = {".pdf", ".md", ".txt", ".json", ".csv", ".tsv", ".xlsx"}
CHUNK_TOKENS = 64
SAMPLE_ROWS = 5
WEB_MAX_PAGES = 50
WEB_TIMEOUT_SECONDS = 15
WEB_MAX_BYTES = 5 * 1024 * 1024
WEB_USER_AGENT = "RAGpipe/1.0 (+local knowledge collection importer)"


def collection_path(name):
    if not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9_-]{0,63}", name):
        raise ValueError("Collection names must use letters, numbers, underscores, or hyphens")
    return data_root() / "collections" / name / "manifest.json"


def is_web_source(value):
    return urlsplit(str(value)).scheme.lower() in {"http", "https"}


def normalize_web_url(value):
    """Drop fragments and common tracking parameters while preserving useful queries."""
    from urllib.parse import parse_qsl, urlencode

    value = urldefrag(str(value).strip())[0]
    parsed = urlsplit(value)
    if parsed.scheme.lower() not in {"http", "https"} or not parsed.netloc:
        raise ValueError("Website imports require a complete http:// or https:// URL")
    query = urlencode([(key, item) for key, item in parse_qsl(parsed.query, keep_blank_values=True)
                       if not key.lower().startswith("utm_") and key.lower() not in {"fbclid", "gclid"}])
    return urlunsplit((parsed.scheme.lower(), parsed.netloc.lower(), parsed.path or "/", query, ""))


def web_sections(html, url):
    """Extract heading-aware source text, excluding common website boilerplate."""
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(html, "html.parser")
    page_title = soup.title.get_text(" ", strip=True) if soup.title else url
    root = soup.find("main") or soup.find("article") or soup.find(attrs={"role": "main"}) or soup.body
    if root is None:
        return []
    for element in root.select("script, style, noscript, nav, header, footer, aside, form, button, svg, canvas, iframe, template"):
        element.decompose()
    for element in root.select('[hidden], [aria-hidden="true"]'):
        element.decompose()

    block_names = ("p", "li", "dt", "dd", "pre", "blockquote", "table")
    sections = []
    heading = page_title
    body = []

    def flush():
        text = "\n".join(body).strip()
        if text:
            sections.append((heading, f"web section: {heading}", text, "passage"))
        body.clear()

    for element in root.find_all(["h1", "h2", "h3", "h4", "h5", "h6", *block_names]):
        if element.name.startswith("h"):
            flush()
            heading = element.get_text(" ", strip=True) or page_title
            continue
        if element.find_parent(block_names):
            continue
        text = " ".join(element.get_text(" ", strip=True).split())
        if text and (not body or text != body[-1]):
            body.append(text)
    flush()
    if not sections:
        text = " ".join(root.get_text(" ", strip=True).split())
        if text:
            sections.append((page_title, f"web page: {page_title}", text, "passage"))
    return sections


def crawl_website(start_url, max_pages=WEB_MAX_PAGES, session=None):
    """Crawl a bounded same-site subtree and return source-preserving page sections."""
    import requests

    if max_pages < 1:
        raise ValueError("max_pages must be at least 1")
    start_url = normalize_web_url(start_url)
    client = session or requests.Session()
    client.headers.update({"User-Agent": WEB_USER_AGENT})
    start = urlsplit(start_url)
    if start.path.endswith("/"):
        scope_root = start.path
    elif not Path(start.path).suffix:
        scope_root = start.path.rstrip("/") + "/"
    else:
        scope_root = start.path.rsplit("/", 1)[0] + "/"

    def in_scope(path):
        return path == start.path or path.startswith(scope_root)
    robots_url = urlunsplit((start.scheme, start.netloc, "/robots.txt", "", ""))
    robots = RobotFileParser()
    robots.set_url(robots_url)
    warnings = []
    try:
        response = client.get(robots_url, timeout=WEB_TIMEOUT_SECONDS)
        robots.parse(response.text.splitlines() if response.ok else [])
    except requests.RequestException as exc:
        robots.parse([])
        warnings.append(f"Could not read robots.txt: {exc}; continuing with the requested site")

    queue = deque([start_url])
    queued = {start_url}
    attempted = set()
    retained = set()
    pages = []
    while queue and len(attempted) < max_pages:
        url = queue.popleft()
        attempted.add(url)
        if not robots.can_fetch(WEB_USER_AGENT, url):
            warnings.append(f"Skipped {url}: blocked by robots.txt")
            continue
        try:
            response = client.get(url, timeout=WEB_TIMEOUT_SECONDS)
            response.raise_for_status()
            final_url = normalize_web_url(response.url)
            final = urlsplit(final_url)
            if final.netloc != start.netloc or not in_scope(final.path):
                warnings.append(f"Skipped redirected page outside website scope: {final_url}")
                continue
            content_type = response.headers.get("Content-Type", "").split(";", 1)[0].lower()
            if content_type not in {"text/html", "application/xhtml+xml"}:
                warnings.append(f"Skipped {final_url}: unsupported content type {content_type or 'unknown'}")
                continue
            if len(response.content) > WEB_MAX_BYTES:
                warnings.append(f"Skipped {final_url}: page exceeds {WEB_MAX_BYTES} bytes")
                continue
            sections = web_sections(response.text, final_url)
            if not sections:
                warnings.append(f"Skipped {final_url}: no readable page content")
                continue
            extracted = "\n".join(f"{title}\n{text}" for title, _, text, _ in sections)
            pages.append((final_url, sha256_text(extracted), sections))
            retained.update({url, final_url})

            from bs4 import BeautifulSoup
            soup = BeautifulSoup(response.text, "html.parser")
            for anchor in soup.find_all("a", href=True):
                try:
                    candidate = normalize_web_url(urljoin(final_url, anchor["href"]))
                except ValueError:
                    continue
                parsed = urlsplit(candidate)
                suffix = Path(parsed.path).suffix.lower()
                if (parsed.netloc == start.netloc and in_scope(parsed.path)
                        and suffix in {"", ".html", ".htm"} and candidate not in queued):
                    queue.append(candidate)
                    queued.add(candidate)
        except requests.RequestException as exc:
            warnings.append(f"Skipped {url}: {exc}; previous records retained")
            retained.add(url)
    complete = not queue
    if queue:
        warnings.append(f"Stopped after {max_pages} pages; increase --max-pages to crawl more")
    return pages, retained, complete, warnings


def legacy_collection():
    payload = read_json(FAQ_JSON_PATH, {})
    faqs = payload.get("faqs", []) if isinstance(payload, dict) else payload
    return {"name": "icer", "enabled": True, "records": [
        {**faq, "collection": "icer", "record_type": "faq",
         "source_id": faq.get("source_id") or sha256_text(str(index)),
         "title": faq["question"], "text": faq["answer"],
         "location": faq.get("section", faq.get("category", "")),
         "imported_at": None}
        for index, faq in enumerate(faqs)
    ]}


def list_collections(include_disabled=False):
    collections = [legacy_collection()]
    root = data_root() / "collections"
    for path in sorted(root.glob("*/manifest.json")):
        collection = read_json(path)
        if include_disabled or collection.get("enabled", True):
            collections.append(collection)
    return [collection for collection in collections if collection["records"]]


def set_enabled(name, enabled):
    if name == "icer":
        raise ValueError("To exclude ICER, select a lab with --collection <name>")
    path = collection_path(name)
    if not path.exists():
        raise ValueError(f"Collection {name!r} does not exist")
    with FileLock(str(path) + ".lock"):
        manifest = read_json(path)
        manifest["enabled"] = enabled
        write_json(path, manifest)


def token_chunks(text, tokenizer, limit=CHUNK_TOKENS):
    """Slice original characters using token offsets, preserving source spelling."""
    encoded = tokenizer(text, add_special_tokens=False, return_offsets_mapping=True,
                        truncation=False, verbose=False)
    offsets = encoded["offset_mapping"]
    start = 0
    boundaries = {match.start() for match in re.finditer(r"(?<=[.!?])\s+(?=[A-Z])|\n\s*\n", text)}
    while start < len(offsets):
        end = min(start + limit, len(offsets))
        if end < len(offsets):
            candidates = [i for i in range(start + 1, end + 1)
                          if offsets[i - 1][1] in boundaries]
            if candidates:
                end = candidates[-1]
        window = offsets[start:end]
        if window:
            fragment = text[window[0][0]:window[-1][1]].strip()
            if fragment:
                yield fragment
        start = end


def markdown_sections(text, filename):
    from markdown_it import MarkdownIt
    lines = text.splitlines(keepends=True)
    headings = [token for token in MarkdownIt().parse(text) if token.type == "heading_open"]
    starts = [(0, filename)]
    for heading in headings:
        start, end = heading.map
        starts.append((start, "".join(lines[start:end]).strip().lstrip("# ")))
    for index, (start, title) in enumerate(starts):
        end = starts[index + 1][0] if index + 1 < len(starts) else len(lines)
        body = "".join(lines[start:end]).strip()
        if body:
            yield title, f"lines {start + 1}-{end}", body, "passage"


def dataset_card(filename, sheet, rows):
    sampled = list(itertools.islice(rows, SAMPLE_ROWS + 1))
    if not sampled:
        return ""
    headers = [str(value) if value is not None else "(unnamed)" for value in sampled[0]]
    lines = [f"Dataset: {filename}", f"Sheet: {sheet}", "Columns: " + ", ".join(headers),
             "Bounded sample only; not a complete dataset analysis."]
    for number, row in enumerate(sampled[1:], start=2):
        cells = [f"{headers[i] if i < len(headers) else 'extra column'}={value}"
                 for i, value in enumerate(row) if value is not None]
        lines.append(f"Sample row {number}: " + "; ".join(cells))
    return "\n".join(lines)


def source_sections(path):
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(path)
        for number, page in enumerate(reader.pages, 1):
            text = page.extract_text() or ""
            if not text.strip():
                raise ValueError(f"Page {number} has no extractable text; OCR is not supported")
            yield path.stem, f"page {number}", text, "passage"
    elif suffix == ".md":
        yield from markdown_sections(path.read_text(encoding="utf-8-sig"), path.stem)
    elif suffix == ".txt":
        text = path.read_text(encoding="utf-8-sig")
        yield path.stem, f"lines 1-{len(text.splitlines())}", text, "passage"
    elif suffix == ".json":
        payload = read_json(path)
        records = payload.get("faqs", []) if isinstance(payload, dict) else payload
        if not isinstance(records, list) or not records:
            raise ValueError("Expected a nonempty FAQ JSON list or an object with a faqs list")
        for index, record in enumerate(records):
            if not isinstance(record, dict) or not record.get("question") or not record.get("answer"):
                raise ValueError(f"FAQ {index + 1} must have a question and answer")
            yield str(record["question"]), f"FAQ {index + 1}", str(record["answer"]), "faq"
    elif suffix in {".csv", ".tsv"}:
        with path.open(newline="", encoding="utf-8-sig") as stream:
            rows = csv.reader(stream, delimiter="\t" if suffix == ".tsv" else ",")
            text = dataset_card(path.name, "table", rows)
            yield path.stem, "header and sample rows 2-6 (at most)", text, "dataset_description"
    elif suffix == ".xlsx":
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        try:
            for sheet in workbook:
                def safe_rows():
                    for row in sheet.iter_rows():
                        yield ["[formula not evaluated]" if cell.data_type == "f" else cell.value for cell in row]
                text = dataset_card(path.name, sheet.title, safe_rows())
                yield f"{path.stem}: {sheet.title}", f"sheet {sheet.title}, header and sample rows 2-6 (at most)", text, "dataset_description"
        finally:
            workbook.close()
    else:
        raise ValueError(f"Unsupported file type: {suffix}")


def import_collection(name, source, tokenizer=None, max_pages=WEB_MAX_PAGES, session=None):
    if name.lower() == "icer":
        raise ValueError("icer is managed by scrape.py; choose a separate lab collection name")
    manifest_path = collection_path(name)
    website = is_web_source(source)
    if website:
        source = normalize_web_url(source)
    else:
        source = Path(source).expanduser().resolve()
        if not source.exists():
            raise FileNotFoundError(source)
    if tokenizer is None:
        from transformers import AutoTokenizer
        from config import BI_ENCODER_MODEL
        from model_setup import retrieval_location
        tokenizer = AutoTokenizer.from_pretrained(retrieval_location(BI_ENCODER_MODEL), local_files_only=True)
    manifest_path.parent.mkdir(parents=True, exist_ok=True)
    with FileLock(str(manifest_path) + ".lock"):
        previous = read_json(manifest_path, {"name": name, "records": [], "sources": {}})
        files = [] if website else (
            sorted(p for p in source.rglob("*") if p.is_file()) if source.is_dir() else [source]
        )
        records = list(previous["records"])
        sources = dict(previous.get("sources", {}))
        messages = []
        now = utc_now_iso()
        parser_kind = "web-sections-v1" if website else "sentence-chunks-v3"
        parser_id = f"{parser_kind}:{CHUNK_TOKENS}:{getattr(tokenizer, 'name_or_path', 'custom')}"

        def replace_source(key, digest, sections, source_kind, source_group=None):
            nonlocal records
            old = sources.get(key, {})
            if old.get("hash") == digest and old.get("parser") == parser_id:
                return
            version = old.get("version", 0) + (old.get("hash") != digest)
            new_records = []
            for section_index, (title, location, text, kind) in enumerate(sections):
                # FAQ answers are canonical responses, not partial document excerpts.
                chunks = [text] if kind == "faq" else token_chunks(
                    text, tokenizer, 160 if kind == "dataset_description" else CHUNK_TOKENS)
                for chunk_index, chunk in enumerate(chunks):
                    if kind == "dataset_description":
                        chunk = "Bounded dataset description/sample, not a full analysis.\n" + chunk
                    source_id = sha256_text(f"{name}\n{key}\n{section_index}\n{chunk_index}")[:24]
                    new_records.append({
                        "source_id": source_id, "collection": name, "record_type": kind,
                        "title": title, "text": chunk, "question": title, "answer": chunk,
                        "category": name, "section": title, "url": key,
                        "location": f"{location}, passage {chunk_index + 1}",
                        "source_path": key, "content_hash": sha256_text(chunk),
                        "version": version, "imported_at": now, "scraped_at": None,
                        "fetched_at": now if source_kind == "website" else None,
                    })
            if not new_records:
                raise ValueError("No searchable text found")
            records = [record for record in records if record.get("source_path") != key]
            records.extend(new_records)
            sources[key] = {"hash": digest, "version": version, "imported_at": now,
                            "parser": parser_id, "kind": source_kind, "group": source_group}

        if website:
            pages, retained, complete, crawl_messages = crawl_website(
                source, max_pages=max_pages, session=session)
            messages.extend(crawl_messages)
            for key, digest, sections in pages:
                try:
                    replace_source(key, digest, sections, "website", source)
                except Exception as exc:
                    messages.append(f"Skipped {key}: {exc}; previous records retained")
            if complete:
                deleted = {key for key, metadata in sources.items()
                           if metadata.get("kind") == "website" and metadata.get("group") == source
                           and key not in retained}
                records = [record for record in records if record.get("source_path") not in deleted]
                sources = {key: value for key, value in sources.items() if key not in deleted}
            if not pages and not any(metadata.get("group") == source for metadata in sources.values()):
                raise ValueError("No readable website pages were imported")
            result = {"name": name, "enabled": previous.get("enabled", True), "records": records,
                      "sources": sources, "warnings": messages}
            write_json(manifest_path, result)
            return result

        for path in files:
            if path.suffix.lower() not in SUPPORTED:
                messages.append(f"Skipped {path.name}: unsupported format")
                continue
            key = str(path)
            try:
                with path.open("rb") as stream:
                    digest = hashlib.file_digest(stream, "sha256").hexdigest()
            except OSError as exc:
                messages.append(f"Skipped {path.name}: {exc}; previous records retained")
                continue
            try:
                replace_source(key, digest, source_sections(path), "file")
            except Exception as exc:
                messages.append(f"Skipped {path.name}: {exc}; previous records retained")
                continue
        # A directory reimport synchronizes deleted files in that directory only.
        if source.is_dir():
            deleted = {key for key in sources if Path(key).is_relative_to(source) and not Path(key).exists()}
            records = [record for record in records if record.get("source_path") not in deleted]
            sources = {key: value for key, value in sources.items() if key not in deleted}
        result = {"name": name, "enabled": previous.get("enabled", True), "records": records,
                  "sources": sources, "warnings": messages}
        write_json(manifest_path, result)
        return result
